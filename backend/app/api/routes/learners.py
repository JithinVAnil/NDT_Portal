from __future__ import annotations

import csv
import io

from fastapi import APIRouter, Depends, File, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import bad_request, require_admin_or_department
from app.models import Course, Department, Learner, LearnerProgress, User
from app.models.base import RoleCode
from app.repositories import serializers as ser
from app.repositories.scope import get_learner_checked, scope_learners
from app.schemas.requests import ImportConfirmRequest, LearnerCreate, LearnerUpdate
from app.services import audit, eligibility, learner_profile, workflow
from app.services.settings_service import QUALIFICATION_LEVELS

router = APIRouter(prefix="/learners", tags=["Learners"])

IMPORT_COLUMNS = ["employee_id", "full_name", "email", "mobile", "designation", "qualification", "department_code",
                  "location", "course_code"]


@router.get("")
def list_learners(
    q: str | None = None, department_id: int | None = None, status: str | None = None,
    user: User = Depends(require_admin_or_department), db: Session = Depends(get_db),
):
    stmt = scope_learners(select(Learner), user).order_by(Learner.learner_code)
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(Learner.full_name.ilike(like), Learner.learner_code.ilike(like),
                              Learner.email.ilike(like), Learner.employee_id.ilike(like)))
    if department_id:
        stmt = stmt.where(Learner.department_id == department_id)
    if status:
        stmt = stmt.where(Learner.status == status)
    learners = db.scalars(stmt).all()
    progress = {}
    for p in db.scalars(select(LearnerProgress).where(LearnerProgress.learner_id.in_([l.id for l in learners]))).all():
        progress.setdefault(p.learner_id, p)
    out = []
    for l in learners:
        d = ser.learner(l)
        p = progress.get(l.id)
        d["course"] = p.course.short_title if p else None
        d["progress"] = p.overall_pct if p else None
        d["attendance"] = p.attendance_pct if p else None
        d["risk_level"] = p.risk_level if p else None
        out.append(d)
    return out


@router.get("/qualifications")
def qualifications(user: User = Depends(require_admin_or_department)):
    return QUALIFICATION_LEVELS


@router.get("/import/template")
def import_template(user: User = Depends(require_admin_or_department), format: str = "csv"):
    dept = user.department.code if user.department else "SCIB"
    sample_rows = [
        [f"{dept}-EMP-5101", "Aditi Rao", "aditi.rao@demo-gov.local", "9876500011", "Assistant Section Officer",
         "Graduate", dept, "New Delhi", "CSA"],
        [f"{dept}-EMP-5102", "Farhan Qureshi", "farhan.q@demo-gov.local", "9876500012", "Technical Officer",
         "Graduate (Engineering/IT)", dept, "New Delhi", "CLF"],
    ]
    if format == "xlsx":
        import io as _io

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Learners"
        ws.append(IMPORT_COLUMNS)
        for r in sample_rows:
            ws.append(r)
        buf = _io.BytesIO()
        wb.save(buf)
        return StreamingResponse(iter([buf.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=learner_import_template.xlsx"})
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(IMPORT_COLUMNS)
    for r in sample_rows:
        w.writerow(r)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=learner_import_template.csv"})


def _parse_upload(filename: str, raw: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip().lower() for h in rows[0]]
        return [{header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header)}
                for r in rows[1:] if any(v not in (None, "") for v in r)]
    if name.endswith(".csv") or name.endswith(".txt"):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in reader
                if any((v or "").strip() for v in r.values())]
    raise bad_request("Unsupported file type. Upload a .csv or .xlsx file")


def _validate_rows(db: Session, user: User, rows: list[dict], default_course_id: int | None) -> dict:
    seen_emails: set[str] = set()
    seen_ids: set[str] = set()
    depts = {d.code: d for d in db.scalars(select(Department)).all()}
    courses = {c.code: c for c in db.scalars(select(Course)).all()}
    default_course = db.get(Course, default_course_id) if default_course_id else None
    out = []
    for i, raw in enumerate(rows, start=2):
        errors: list[str] = []
        row = {k: raw.get(k, "") for k in IMPORT_COLUMNS}
        dept_code = (row["department_code"] or "").upper()
        if user.role_code == RoleCode.DEPARTMENT:
            dept = user.department
            if dept_code and dept_code != dept.code:
                errors.append(f"Department {dept_code} does not match your department ({dept.code})")
        else:
            dept = depts.get(dept_code)
            if not dept:
                errors.append(f"Unknown department code '{dept_code}'")
        payload = {k: row[k] or None for k in ("employee_id", "full_name", "email", "mobile", "designation",
                                                "qualification", "location")}
        clean = None
        try:
            clean = LearnerCreate(**payload, department_id=dept.id if dept else None)
        except ValidationError as exc:
            for e in exc.errors():
                errors.append(f"{e['loc'][0] if e['loc'] else 'row'}: {e['msg'].replace('Value error, ', '')}")
        if clean and clean.qualification and clean.qualification not in QUALIFICATION_LEVELS:
            errors.append(f"Qualification '{clean.qualification}' is not a recognised level")
        course = courses.get((row["course_code"] or "").upper()) or default_course
        if row["course_code"] and not courses.get(row["course_code"].upper()):
            errors.append(f"Unknown course code '{row['course_code']}'")

        duplicate = False
        dup_reason = None
        if clean:
            if clean.email in seen_emails or clean.employee_id in seen_ids:
                duplicate, dup_reason = True, "Duplicate row within file"
            elif db.scalar(select(Learner.id).where(func.lower(Learner.email) == clean.email)):
                duplicate, dup_reason = True, "Email already registered"
            elif db.scalar(select(Learner.id).where(Learner.employee_id == clean.employee_id)):
                duplicate, dup_reason = True, "Employee ID already registered"
            seen_emails.add(clean.email)
            seen_ids.add(clean.employee_id)

        elig = None
        elig_messages: list[str] = []
        if clean and not errors and course and dept:
            temp = Learner(learner_code="PREVIEW", employee_id=clean.employee_id, full_name=clean.full_name,
                           email=clean.email, designation=clean.designation, qualification=clean.qualification,
                           department_id=dept.id)
            elig, rules = eligibility.evaluate(db, temp, course, dept)
            elig_messages = [r["message"] for r in rules if r["outcome"] != "PASS"]
        status = "Invalid" if errors else ("Duplicate" if duplicate else "Valid")
        out.append({
            "row": i, "data": row, "status": status, "errors": errors, "duplicate_reason": dup_reason,
            "course_code": course.code if course else None, "eligibility": elig, "eligibility_notes": elig_messages,
        })
    summary = {
        "total": len(out),
        "valid": sum(1 for r in out if r["status"] == "Valid"),
        "invalid": sum(1 for r in out if r["status"] == "Invalid"),
        "duplicate": sum(1 for r in out if r["status"] == "Duplicate"),
        "eligible": sum(1 for r in out if r["status"] == "Valid" and r["eligibility"] == "ELIGIBLE"),
        "review_required": sum(1 for r in out if r["status"] == "Valid" and r["eligibility"] == "REVIEW_REQUIRED"),
        "ineligible": sum(1 for r in out if r["status"] == "Valid" and r["eligibility"] == "NOT_ELIGIBLE"),
    }
    return {"rows": out, "summary": summary}


@router.post("/import/preview")
async def import_preview(request: Request, file: UploadFile = File(...), course_id: int | None = None,
                         user: User = Depends(require_admin_or_department), db: Session = Depends(get_db)):
    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise bad_request("File too large (max 5 MB)")
    rows = _parse_upload(file.filename or "", raw)
    if not rows:
        raise bad_request("No data rows found in the file")
    if len(rows) > 1000:
        raise bad_request("Maximum 1000 rows per import")
    result = _validate_rows(db, user, rows, course_id)
    result["filename"] = file.filename
    audit.record(db, user, "BULK_IMPORT_PREVIEW", "Learner", None,
                 details=f"{file.filename}: {result['summary']}", request=request)
    db.commit()
    return result


@router.post("/import/confirm")
def import_confirm(body: ImportConfirmRequest, request: Request, user: User = Depends(require_admin_or_department),
                   db: Session = Depends(get_db)):
    validated = _validate_rows(db, user, body.rows, body.course_id)
    created, nominations = [], []
    for r in validated["rows"]:
        if r["status"] != "Valid":
            continue
        d = r["data"]
        dept = user.department if user.role_code == RoleCode.DEPARTMENT else db.scalar(
            select(Department).where(Department.code == d["department_code"].upper()))
        clean = LearnerCreate(**{k: d[k] or None for k in ("employee_id", "full_name", "email", "mobile", "designation",
                                                            "qualification", "location")}, department_id=dept.id)
        learner = Learner(learner_code=workflow.next_code(db, Learner, Learner.learner_code, "NL"),
                          **clean.model_dump(exclude={"department_id"}), department_id=dept.id, created_by_id=user.id)
        db.add(learner)
        db.flush()
        created.append(learner.learner_code)
        if body.create_nominations and r["course_code"]:
            course = db.scalar(select(Course).where(Course.code == r["course_code"]))
            nom = workflow.create_nomination(db, user, learner, course, body.training_mode,
                                             "Bulk import nomination", body.submit, request)
            nominations.append(nom.nomination_no)
    audit.record(db, user, "BULK_IMPORT", "Learner", None,
                 details=f"{len(created)} learners, {len(nominations)} nominations imported", request=request)
    db.commit()
    return {"created": len(created), "learner_codes": created, "nominations": nominations,
            "skipped": validated["summary"]["total"] - len(created), "summary": validated["summary"]}


@router.post("")
def create_learner(body: LearnerCreate, request: Request, user: User = Depends(require_admin_or_department),
                   db: Session = Depends(get_db)):
    learner = create_learner_record(db, user, body)
    audit.record(db, user, "LEARNER_CREATED", "Learner", learner.learner_code, details=learner.full_name, request=request)
    db.commit()
    return ser.learner(learner)


def create_learner_record(db: Session, user: User, body: LearnerCreate) -> Learner:
    dept_id = user.department_id if user.role_code == RoleCode.DEPARTMENT else body.department_id
    if not dept_id or not db.get(Department, dept_id):
        raise bad_request("Select a valid department")
    if db.scalar(select(Learner.id).where(func.lower(Learner.email) == body.email)):
        raise bad_request("A learner with this email already exists")
    if db.scalar(select(Learner.id).where(Learner.employee_id == body.employee_id)):
        raise bad_request("A learner with this employee ID already exists")
    learner = Learner(learner_code=workflow.next_code(db, Learner, Learner.learner_code, "NL"),
                      **body.model_dump(exclude={"department_id"}), department_id=dept_id, created_by_id=user.id)
    db.add(learner)
    db.flush()
    return learner


@router.get("/{learner_id}")
def get_learner(learner_id: int, user: User = Depends(require_admin_or_department), db: Session = Depends(get_db)):
    learner = get_learner_checked(db, user, learner_id)
    return learner_profile.full_profile(db, learner)


@router.put("/{learner_id}")
def update_learner(learner_id: int, body: LearnerUpdate, request: Request,
                   user: User = Depends(require_admin_or_department), db: Session = Depends(get_db)):
    learner = get_learner_checked(db, user, learner_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(learner, k, v)
    audit.record(db, user, "LEARNER_UPDATED", "Learner", learner.learner_code, request=request)
    db.commit()
    return ser.learner(learner)
