"""Free, local report exporters: CSV (csv), Excel (openpyxl), PDF (ReportLab)."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone, timedelta

DEMO_BANNER = "DEMONSTRATION ENVIRONMENT — SAMPLE DATA"
ORG = "NIELIT · Digital Training & Forensics Capacity Building Programme (Prototype Portal)"
IST = timezone(timedelta(hours=5, minutes=30))


def _stamp() -> str:
    return datetime.now(IST).strftime("%d %b %Y %H:%M IST")


def _cell(v):
    if v is None:
        return ""
    return v


def to_csv(title: str, columns: list[tuple[str, str]], rows: list[dict], generated_by: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"{title} — {DEMO_BANNER}"])
    w.writerow([f"Generated {_stamp()} by {generated_by}"])
    w.writerow([])
    w.writerow([label for _, label in columns])
    for r in rows:
        w.writerow([_cell(r.get(k)) for k, _ in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(title: str, columns: list[tuple[str, str]], rows: list[dict], generated_by: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    navy = PatternFill("solid", fgColor="0B2545")
    amber = PatternFill("solid", fgColor="FFF4D6")
    thin = Side(style="thin", color="D0D7E2")
    n = len(columns)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color="0B2545")
    ws.cell(row=2, column=1, value=ORG).font = Font(size=9, color="5A6475")
    c = ws.cell(row=3, column=1, value=DEMO_BANNER)
    c.font = Font(bold=True, size=9, color="8A5A00")
    c.fill = amber
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(n, 1))
    ws.cell(row=4, column=1, value=f"Generated {_stamp()} by {generated_by} · {len(rows)} rows").font = Font(size=9, color="5A6475")
    hdr = 6
    for i, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=hdr, column=i, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = navy
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    widths = [len(label) for _, label in columns]
    for r_i, r in enumerate(rows, start=hdr + 1):
        for c_i, (k, _) in enumerate(columns, start=1):
            v = _cell(r.get(k))
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if r_i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F7FA")
            widths[c_i - 1] = max(widths[c_i - 1], min(len(str(v)), 60))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, w + 2)
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(n)}{hdr + len(rows)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(title: str, columns: list[tuple[str, str]], rows: list[dict], generated_by: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=16 * mm,
                            bottomMargin=14 * mm, title=title, author="NIELIT Prototype Portal")
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7, leading=8.5)
    head_style = ParagraphStyle("head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")

    def safe(v) -> str:
        s = str(_cell(v))
        s = s.replace("✓", "Yes").replace("✕", "No").replace("→", "->").replace("—", "-").replace("–", "-")
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [
        Paragraph(f"<b>{safe(title)}</b>", ParagraphStyle("t", parent=styles["Title"], fontSize=15, alignment=0,
                                                           textColor=colors.HexColor("#0B2545"), spaceAfter=2)),
        Paragraph(safe(ORG), ParagraphStyle("o", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#5A6475"))),
        Paragraph(f"Generated {safe(_stamp())} by {safe(generated_by)} - {len(rows)} rows",
                  ParagraphStyle("g", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#5A6475"))),
        Spacer(1, 4 * mm),
    ]
    data = [[Paragraph(safe(label), head_style) for _, label in columns]]
    for r in rows:
        data.append([Paragraph(safe(r.get(k)), cell_style) for k, _ in columns])
    if not rows:
        data.append([Paragraph("No records found for the selected filters.", cell_style)] + [""] * (len(columns) - 1))
    avail = page[0] - 24 * mm
    weights = []
    for k, label in columns:
        longest = max([max(len(w) for w in label.split())] + [len(str(_cell(r.get(k)))) for r in rows[:200]])
        weights.append(min(max(longest, 5), 42) + 4)  # +4 ≈ cell padding, keeps short IDs on one line
    total = sum(weights) or 1
    col_widths = [avail * w / total for w in weights]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    story.append(table)

    def decorate(canvas, doc_):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#FFF4D6"))
        canvas.rect(0, page[1] - 9 * mm, page[0], 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#8A5A00"))
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawCentredString(page[0] / 2, page[1] - 6 * mm, "DEMONSTRATION ENVIRONMENT - SAMPLE DATA - NOT AN OFFICIAL NIELIT RECORD")
        canvas.setFillColor(colors.HexColor("#5A6475"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(12 * mm, 7 * mm, "NIELIT Digital Training Portal - Prototype")
        canvas.drawRightString(page[0] - 12 * mm, 7 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


EXPORTERS = {
    "csv": (to_csv, "text/csv; charset=utf-8", "csv"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": (to_pdf, "application/pdf", "pdf"),
}
